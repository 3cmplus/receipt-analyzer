# -*- coding: utf-8 -*-
import io, json, re
from flask import Flask, request, jsonify, render_template, send_file
from google import genai
from google.genai import types
from config import GEMINI_API_KEY, GEMINI_MODEL
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

app = Flask(__name__)
client = genai.Client(api_key=GEMINI_API_KEY)

PROMPT = """이 영수증 이미지를 꼼꼼히 분석해서 아래 JSON 형식으로만 응답해줘.
JSON 외에 다른 텍스트, 설명, 마크다운 코드블록은 절대 포함하지 마.

{
  "가맹점명": "상호명",
  "주소": "가맹점 주소 (없으면 확인불가)",
  "결제일시": "YYYY-MM-DD HH:MM",
  "총금액": "숫자만 (예: 15000)",
  "항목": [
    {"품목명": "상품명", "수량": "1", "단가": "숫자만"}
  ]
}

주의사항:
- 한글 상호명·주소·품목명은 그대로 유지해줘
- 읽을 수 없는 항목은 "확인불가"로 표시
- 항목이 없으면 빈 배열 [] 반환
- 금액은 쉼표 없이 숫자만 반환 (예: 12000)
"""

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/analyze", methods=["POST"])
def analyze():
    if "image" not in request.files:
        return jsonify({"error": "이미지 파일이 없습니다."}), 400

    file = request.files["image"]
    if file.filename == "":
        return jsonify({"error": "파일을 선택해주세요."}), 400

    mime_type = file.content_type or "image/jpeg"
    image_bytes = file.read()

    try:
        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=[
                types.Part.from_bytes(data=image_bytes, mime_type=mime_type),
                PROMPT,
            ],
        )
        text = response.text.strip()

        # 혹시 붙는 마크다운 코드블록 제거
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)

        # JSON 객체 추출
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if match:
            text = match.group()

        result = json.loads(text)
        return jsonify(result)

    except json.JSONDecodeError:
        return jsonify({"error": "Gemini 응답을 파싱하지 못했습니다.", "raw": text}), 500
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/save-excel", methods=["POST"])
def save_excel():
    data = request.get_json(force=True)
    if not data:
        return jsonify({"error": "데이터가 없습니다."}), 400
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "영수증 내역"

        thin = Side(style="thin", color="D0D0D0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        hdr_fill = PatternFill("solid", fgColor="1E3A5F")
        hdr_font = Font(bold=True, color="FFFFFF", size=11, name="맑은 고딕")

        headers = ["결제일시", "가맹점명", "주소", "총금액"]
        widths  = [20, 24, 34, 14]
        for i, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(1, i, h)
            cell.fill = hdr_fill; cell.font = hdr_font; cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[cell.column_letter].width = w
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"

        total = data.get("총금액", "")
        try:
            total = int(total)
        except (ValueError, TypeError):
            pass
        ws.append([data.get("결제일시",""), data.get("가맹점명",""),
                   data.get("주소",""), total])
        ws.cell(2, 4).number_format = "#,##0"
        ws.cell(2, 4).alignment = Alignment(horizontal="right")
        for col in range(1, 5):
            ws.cell(2, col).border = border
            ws.cell(2, col).font = Font(name="맑은 고딕", size=10)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="영수증_내역.xlsx",
        )
    except Exception as e:
        return jsonify({"error": str(e)}), 500


if __name__ == "__main__":
    app.run(debug=True, port=5000)
