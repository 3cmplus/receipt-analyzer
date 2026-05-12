# -*- coding: utf-8 -*-
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EXCEL_PATH = Path(__file__).parent / "영수증_내역.xlsx"
HEADERS    = ["결제일시", "가맹점명", "주소", "총금액"]
COL_WIDTHS = [20, 24, 34, 14]


def _thin_border():
    s = Side(style="thin", color="D0D0D0")
    return Border(left=s, right=s, top=s, bottom=s)


def _init_sheet(ws):
    ws.title = "영수증 내역"
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF", size=11,
                       name="맑은 고딕")
    for i, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.fill      = header_fill
        cell.font      = header_font
        cell.alignment = Alignment(horizontal="center",
                                    vertical="center")
        cell.border    = _thin_border()
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"


def append_row(data: dict) -> str:
    """결과 dict 를 엑셀에 한 행 추가하고 저장 경로를 반환."""
    if EXCEL_PATH.exists():
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        _init_sheet(ws)

    total = data.get("총금액", "")
    try:
        total = int(total)
    except (ValueError, TypeError):
        pass

    row_data = [
        data.get("결제일시", ""),
        data.get("가맹점명", ""),
        data.get("주소", ""),
        total,
    ]
    ws.append(row_data)

    r = ws.max_row
    even_fill = PatternFill("solid", fgColor="F2F6FB")
    data_font = Font(name="맑은 고딕", size=10)

    for col in range(1, 5):
        cell = ws.cell(r, col)
        cell.font   = data_font
        cell.border = _thin_border()
        if r % 2 == 0:
            cell.fill = even_fill

    # 총금액 오른쪽 정렬 + 천단위 포맷
    ws.cell(r, 4).number_format = "#,##0"
    ws.cell(r, 4).alignment = Alignment(horizontal="right",
                                          vertical="center")

    wb.save(EXCEL_PATH)
    return str(EXCEL_PATH)
